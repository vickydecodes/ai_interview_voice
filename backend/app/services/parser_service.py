import io
import re
import logging
from typing import Dict, Any
from fastapi import UploadFile, HTTPException, status
import pdfplumber
from pypdf import PdfReader
import docx

from app.models.resume import ParsedProfile, SkillItem, ProjectItem, EducationItem
from app.adapters.llm_adapter import llm_adapter

logger = logging.getLogger(__name__)

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def sanitize_text(text: str) -> str:
    """
    Sanitize extracted raw resume text by removing null characters,
    unprintable controls, and excess whitespace while preserving structure.
    """
    # Replace control characters except tab and newline
    sanitized = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]", "", text)
    # Normalize excessive newlines
    sanitized = re.sub(r"\n{3,}", "\n\n", sanitized)
    return sanitized.strip()


async def extract_text_from_file(file: UploadFile) -> str:
    filename = file.filename or ""
    ext = ""
    if "." in filename:
        ext = "." + filename.rsplit(".", 1)[-1].lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Allowed formats: PDF, DOCX, TXT.",
        )

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File exceeds maximum allowed size of 5 MB ({len(content)} bytes).",
        )

    extracted_text = ""

    try:
        if ext == ".pdf":
            # Attempt with pdfplumber first
            try:
                with pdfplumber.open(io.BytesIO(content)) as pdf:
                    pages_text = [page.extract_text() or "" for page in pdf.pages]
                    extracted_text = "\n".join(pages_text)
            except Exception as e:
                logger.warning(f"pdfplumber extraction failed, falling back to pypdf: {e}")
                reader = PdfReader(io.BytesIO(content))
                pages_text = [p.extract_text() or "" for p in reader.pages]
                extracted_text = "\n".join(pages_text)

        elif ext == ".docx":
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    paragraphs.extend([cell.text for cell in row.cells if cell.text])
            extracted_text = "\n".join(paragraphs)

        elif ext == ".txt":
            try:
                extracted_text = content.decode("utf-8")
            except UnicodeDecodeError:
                extracted_text = content.decode("latin-1", errors="ignore")

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error parsing file '{filename}': {e}")
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Unable to read text from file: {str(e)}",
        )

    sanitized = sanitize_text(extracted_text)
    if not sanitized or len(sanitized) < 30:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded resume does not contain enough extractable text. Please ensure it is not a scanned image.",
        )

    # Note: resume content is NEVER logged as per security guidelines
    logger.info(f"Successfully extracted {len(sanitized)} characters from '{filename}'")
    return sanitized


async def parse_resume_with_llm(raw_text: str) -> ParsedProfile:
    """
    Passes sanitized resume text to the LLM with a strict JSON schema prompt.
    """
    system_prompt = (
        "You are an expert technical talent evaluator and parser. "
        "Analyze the provided resume and return a strict JSON object conforming to this schema:\n"
        "{\n"
        '  "full_name": "string (candidate name if found)",\n'
        '  "seniority": "Junior | Mid-level | Senior | Staff/Lead",\n'
        '  "years_of_experience": number (e.g. 3.5),\n'
        '  "roles": ["string", ...],\n'
        '  "domains": ["string", ...],\n'
        '  "skills": [{"name": "string", "proficiency": 1 to 5, "category": "Languages | Frontend | Backend | Databases | Cloud & DevOps | Tools | Architecture | AI/ML"}],\n'
        '  "projects": [{"title": "string", "description": "string", "technologies": ["string"]}],\n'
        '  "education": [{"degree": "string", "institution": "string", "year": "string"}],\n'
        '  "strengths": ["string", ...],\n'
        '  "weaknesses": ["string", ...]\n'
        "}\n"
        "Guidelines:\n"
        "- Rate skill proficiency honestly from 1 (beginner) to 5 (distinguished/expert) based on usage and depth.\n"
        "- Identify 2 to 4 tangible strengths and 2 to 3 constructive skill/experience gaps or weaknesses."
    )

    user_prompt = f"Please extract the candidate profile from the following resume text:\n\n{raw_text[:12000]}"

    parsed_json = await llm_adapter.generate_json(prompt=user_prompt, system_prompt=system_prompt)

    # Convert to typed ParsedProfile
    skills = [
        SkillItem(
            name=s.get("name", "Unknown"),
            proficiency=int(s.get("proficiency", 3)),
            category=s.get("category", "General"),
        )
        for s in parsed_json.get("skills", [])
        if isinstance(s, dict) and s.get("name")
    ]

    projects = [
        ProjectItem(
            title=p.get("title", "Project"),
            description=p.get("description", ""),
            technologies=p.get("technologies", []),
        )
        for p in parsed_json.get("projects", [])
        if isinstance(p, dict)
    ]

    education = [
        EducationItem(
            degree=e.get("degree", "Degree"),
            institution=e.get("institution", "Institution"),
            year=str(e.get("year", "")),
        )
        for e in parsed_json.get("education", [])
        if isinstance(e, dict)
    ]

    return ParsedProfile(
        full_name=parsed_json.get("full_name") or "Candidate",
        seniority=parsed_json.get("seniority") or "Mid-level",
        years_of_experience=float(parsed_json.get("years_of_experience") or 2.0),
        roles=parsed_json.get("roles") or ["Software Engineer"],
        domains=parsed_json.get("domains") or ["Web Development"],
        skills=skills,
        projects=projects,
        education=education,
        strengths=parsed_json.get("strengths") or [],
        weaknesses=parsed_json.get("weaknesses") or [],
    )
